from __future__ import annotations

import copy
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest


SCRIPT = Path(__file__).parents[1] / "scripts" / "evidence_package.py"
sys.path.insert(0, str(SCRIPT.parent))
SPEC = importlib.util.spec_from_file_location("evidence_package", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
SPEC.loader.exec_module(MODULE)

BUILD_SCRIPT = SCRIPT.parent / "build_office.py"
BUILD_SPEC = importlib.util.spec_from_file_location("build_office_for_package_test", BUILD_SCRIPT)
BUILD = importlib.util.module_from_spec(BUILD_SPEC)
assert BUILD_SPEC and BUILD_SPEC.loader
BUILD_SPEC.loader.exec_module(BUILD)

VALIDATE_SCRIPT = SCRIPT.parent / "validate_docx.py"
VALIDATE_SPEC = importlib.util.spec_from_file_location("validate_docx_for_package_test", VALIDATE_SCRIPT)
VALIDATE = importlib.util.module_from_spec(VALIDATE_SPEC)
assert VALIDATE_SPEC and VALIDATE_SPEC.loader
VALIDATE_SPEC.loader.exec_module(VALIDATE)


def render_dependencies_available() -> bool:
    try:
        import docx  # noqa: F401
        import pptx  # noqa: F401
        import reportlab  # noqa: F401
        return True
    except ImportError:
        return False


VALID = {
    "schema_version": "oks-office-evidence/v1",
    "request": {"title": "中文申报书", "audience": "评审专家", "deliverables": ["docx", "pdf", "pptx"]},
    "recall": {"query": "agent 技术", "items": [{"slug": "agent-page", "score": 0.9}]},
    "summary": {"text": "证据边界内的摘要。", "claim_refs": ["claim-1"]},
    "claims": [
        {
            "id": "claim-1",
            "text": "系统支持可追溯输出。",
            "source_refs": ["src-1"],
            "confidence": "high",
            "review_status": "reviewed",
        },
        {
            "id": "claim-2",
            "text": "指标为 0，开关为 False。",
            "source_refs": ["src-1"],
            "confidence": "medium",
            "review_status": "provisional",
        },
    ],
    "sections": [
        {
            "id": "section-1",
            "title": "能力",
            "claim_refs": ["claim-1", "claim-2"],
            "blocks": [
                {"type": "paragraph", "text": "能力说明。", "claim_refs": ["claim-1"]},
                {
                    "type": "bullets",
                    "items": [{"text": "证据链保留。", "claim_refs": ["claim-1"]}, "指标保真。"],
                    "claim_refs": ["claim-1", "claim-2"],
                },
                {"type": "table", "rows": [["字段", "值"], ["count", 0], ["enabled", False]], "claim_refs": ["claim-2"]},
            ],
        }
    ],
    "sources": [{"id": "src-1", "kind": "wiki", "label": "Wiki: agent-page", "locator": "wiki/agent.md", "status": "reviewed"}],
}


class EvidencePackageTests(unittest.TestCase):
    def test_claim_sources_are_required_and_known(self):
        package = copy.deepcopy(VALID)
        del package["claims"][0]["source_refs"]
        with self.assertRaisesRegex(ValueError, "source_refs"):
            MODULE.validate_package(package)

        package = copy.deepcopy(VALID)
        package["claims"][0]["source_refs"] = ["missing"]
        with self.assertRaisesRegex(ValueError, "unknown ids"):
            MODULE.validate_package(package)

    def test_external_research_requires_url_and_retrieval_time(self):
        package = copy.deepcopy(VALID)
        package["sources"][0] = {
            "id": "src-1",
            "kind": "research",
            "label": "Research: example",
            "locator": "records/research/example.md",
            "status": "unverified",
        }
        with self.assertRaisesRegex(ValueError, "http\\(s\\) URL"):
            MODULE.validate_package(package)

        package["sources"][0]["locator"] = "https://example.com/report"
        with self.assertRaisesRegex(ValueError, "retrieved_at"):
            MODULE.validate_package(package)

        package["sources"][0]["retrieved_at"] = "2026-08-27T10:00:00+08:00"
        MODULE.validate_package(package)

    def test_unknown_source_kind_fails_closed(self):
        package = copy.deepcopy(VALID)
        package["sources"][0]["kind"] = "search_snippet"
        with self.assertRaisesRegex(ValueError, "kind must be one of"):
            MODULE.validate_package(package)

    def test_normalization_retains_external_research_metadata(self):
        package = copy.deepcopy(VALID)
        package["sources"][0] = {
            "id": "src-1",
            "kind": "research",
            "label": "Research: issuer report",
            "locator": "https://example.com/report",
            "retrieved_at": "2026-08-27T10:00:00+08:00",
            "status": "unverified",
        }
        package["claims"][0]["review_status"] = "provisional"
        outline = MODULE.package_to_outline(package)
        source = outline["sources"][0]
        self.assertEqual(source["path"], "https://example.com/report")
        self.assertEqual(source["kind"], "research")
        self.assertEqual(source["retrieved_at"], "2026-08-27T10:00:00+08:00")

    def test_external_research_fixture_is_provisional_and_traceable(self):
        fixture = Path(__file__).parent / "fixtures" / "external-research-evidence-package.json"
        package = MODULE.load_package(fixture)
        self.assertEqual({source["status"] for source in package["sources"]}, {"unverified"})
        self.assertEqual({claim["review_status"] for claim in package["claims"]}, {"provisional"})
        self.assertTrue(all(source["locator"].startswith("https://") for source in package["sources"]))

    def test_block_claims_cannot_escape_section_scope(self):
        package = copy.deepcopy(VALID)
        package["sections"][0]["claim_refs"] = ["claim-1"]
        package["sections"][0]["blocks"][0]["claim_refs"] = ["claim-2"]
        with self.assertRaisesRegex(ValueError, "subset"):
            MODULE.validate_package(package)

    def test_irregular_tables_fail_closed(self):
        package = copy.deepcopy(VALID)
        package["sections"][0]["blocks"][2]["rows"].append(["extra", "cell", "rejected"])
        with self.assertRaisesRegex(ValueError, "rectangular"):
            MODULE.validate_package(package)

    def test_normalization_preserves_claim_source_mapping_and_scalars(self):
        outline = MODULE.package_to_outline(copy.deepcopy(VALID))
        self.assertEqual(outline["title"], "中文申报书")
        self.assertEqual(outline["summary_source_refs"], ["src-1"])
        self.assertEqual(outline["sections"][0]["source_refs"], ["src-1"])
        self.assertEqual(outline["sections"][0]["tables"][0]["rows"][1], ["count", 0])
        self.assertEqual(outline["sections"][0]["tables"][0]["rows"][2], ["enabled", False])
        self.assertEqual(outline["sources"][0]["path"], "wiki/agent.md")

    def test_multiple_tables_keep_specific_sources_and_claim_statuses(self):
        fixture = Path(__file__).parent / "fixtures" / "evidence-package.json"
        outline = MODULE.package_to_outline(MODULE.load_package(fixture))
        tables = outline["sections"][0]["tables"]
        self.assertEqual(len(tables), 2)
        self.assertEqual(tables[0]["source_refs"], ["src-1"])
        self.assertEqual(tables[1]["source_refs"], ["src-2"])
        self.assertEqual(tables[0]["claim_statuses"], ["provisional"])
        self.assertEqual(tables[1]["claim_statuses"], ["reviewed"])

    def test_cli_fixture_round_trip(self):
        with tempfile.TemporaryDirectory() as directory:
            package_path = Path(directory) / "package.json"
            outline_path = Path(directory) / "outline.json"
            package_path.write_text(json.dumps(VALID, ensure_ascii=False), encoding="utf-8")
            loaded = MODULE.load_package(package_path)
            MODULE.package_to_outline_file(package_path, outline_path)
            self.assertEqual(json.loads(outline_path.read_text(encoding="utf-8"))["title"], loaded["request"]["title"])

    @unittest.skipUnless(render_dependencies_available(), "optional Office renderer dependencies unavailable")
    def test_all_renderers_read_back_table_sources_and_statuses(self):
        package = MODULE.load_package(Path(__file__).parent / "fixtures" / "evidence-package.json")
        package = copy.deepcopy(package)
        package["request"]["title"] = "Office evidence round trip"
        package["request"]["audience"] = "Reviewers"
        package["summary"]["text"] = "One package, three adapters."
        package["claims"][0]["text"] = "Claim one remains traceable."
        package["claims"][1]["text"] = "Claim two keeps scalar values."
        package["sections"][0]["title"] = "Evidence fidelity"
        package["sections"][0]["blocks"][0]["text"] = "The normalizer preserves claim text."
        package["sections"][0]["blocks"][1]["items"] = ["Source status stays visible", "Long text is chunked safely"]
        package["sections"][0]["blocks"][2]["rows"] = [["field", "value"], ["count", 0], ["enabled", False]]
        package["sections"][0]["blocks"][3]["rows"] = [["status", "meaning"], ["review", "kept"]]
        package["sources"][0]["label"] = "Research: office-output"
        package["sources"][1]["label"] = "Wiki: office-page"

        outline = MODULE.package_to_outline(package)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            outputs = {
                "docx": root / "out.docx",
                "pdf": root / "out.pdf",
                "pptx": root / "out.pptx",
            }
            BUILD.render_docx(outline, outputs["docx"])
            BUILD.render_pdf(outline, outputs["pdf"], None)
            BUILD.render_pptx(outline, outputs["pptx"])
            from docx import Document
            from pypdf import PdfReader
            from pptx import Presentation

            docx_structure = VALIDATE.validate_docx(outputs["docx"])
            self.assertTrue(docx_structure["valid"], docx_structure)
            self.assertEqual(docx_structure["tables"], 2)
            docx_text = "\n".join(paragraph.text for paragraph in Document(outputs["docx"]).paragraphs)
            pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(outputs["pdf"]).pages)
            pptx_text = "\n".join(
                shape.text for slide in Presentation(outputs["pptx"]).slides for shape in slide.shapes if hasattr(shape, "text")
            )
            for artifact_text in (docx_text, pdf_text, pptx_text):
                self.assertIn("src-1", artifact_text)
                self.assertIn("src-2", artifact_text)
                self.assertIn("provisional", artifact_text)
                self.assertIn("reviewed", artifact_text)


if __name__ == "__main__":
    unittest.main()
