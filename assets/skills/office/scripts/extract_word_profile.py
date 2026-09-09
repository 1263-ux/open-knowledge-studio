#!/usr/bin/env python3
"""Extract a portable OKS Word design profile from a local DOCX template."""
from __future__ import annotations

import argparse
import copy
import json
from pathlib import Path
import sys

SCRIPT_DIR = Path(__file__).parent
sys.path.insert(0, str(SCRIPT_DIR))
from build_office import DEFAULT_WORD_PROFILE  # noqa: E402


def _east_asia_font(style) -> str | None:
    rpr = style._element.rPr
    if rpr is None:
        return None
    fonts = rpr.rFonts
    return fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia") if fonts is not None else None


def _hex_color(value) -> str | None:
    if value is None:
        return None
    return value.upper().lstrip("#")


def extract_profile(template: Path) -> dict:
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(template)
    profile = copy.deepcopy(DEFAULT_WORD_PROFILE)
    section = document.sections[0]
    profile["page"] = {
        "top_twips": int(section.top_margin.twips),
        "bottom_twips": int(section.bottom_margin.twips),
        "left_twips": int(section.left_margin.twips),
        "right_twips": int(section.right_margin.twips),
    }
    normal = document.styles["Normal"]
    if normal.font.name:
        profile["fonts"]["latin"] = normal.font.name
    east_asia = _east_asia_font(normal)
    if east_asia:
        profile["fonts"]["east_asia"] = east_asia
    if normal.font.size:
        profile["fonts"]["body_pt"] = float(normal.font.size.pt)
    if document.tables:
        table = document.tables[0]
        widths = []
        for cell in table.rows[0].cells:
            tc_w = cell._tc.tcPr.tcW if cell._tc.tcPr is not None else None
            width = int(tc_w.get(qn("w:w"))) if tc_w is not None and tc_w.get(qn("w:w")) else 0
            widths.append(width)
        if widths and sum(widths) > 0:
            total = sum(widths)
            profile["table"]["width_twips"] = total
            profile["table"]["column_widths"] = [round(width / total, 4) for width in widths]
        first_cell = table.rows[0].cells[0]
        shading = first_cell._tc.tcPr.find(qn("w:shd")) if first_cell._tc.tcPr is not None else None
        fill = _hex_color(shading.get(qn("w:fill"))) if shading is not None else None
        if fill and fill not in {"AUTO", "FFFFFF"}:
            profile["colors"]["accent"] = fill
    profile["source"] = {"kind": "template-profile", "template_name": template.name}
    return profile


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    if not args.template.is_file():
        raise SystemExit(f"template does not exist: {args.template}")
    profile = extract_profile(args.template)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(profile, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
