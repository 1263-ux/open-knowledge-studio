#!/usr/bin/env python3
"""Render one curated OKS outline or evidence package as DOCX, PDF, or PPTX.

This is deliberately a renderer, not a knowledge selector. Recall, evidence
review, and source-ledger construction remain in the office skill workflow.
"""
from __future__ import annotations

import argparse
import copy
from contextlib import contextmanager
import html
import json
import os
from pathlib import Path
import tempfile
from typing import Any

from evidence_package import load_package, package_to_outline


VALID_SOURCE_STATUSES = {
    "reviewed",
    "partial",
    "failed",
    "skipped",
    "environment_limited",
    "unverified",
    "synthesis",
}

DEFAULT_WORD_PROFILE: dict[str, Any] = {
    "profile_version": "oks-word/v1",
    "page": {"top_twips": 1008, "bottom_twips": 1008, "left_twips": 1152, "right_twips": 1152},
    "fonts": {"latin": "Aptos", "east_asia": "Microsoft YaHei", "body_pt": 10.5, "table_pt": 9.5},
    "colors": {
        "ink": "0F172A",
        "muted": "64748B",
        "accent": "1F4E78",
        "header_text": "FFFFFF",
        "row_alt": "F1F5F9",
        "border": "B8C7D9",
    },
    "table": {
        "width_twips": 9600,
        "column_widths": [0.30, 0.70],
        "cell_margin_twips": {"top": 90, "start": 120, "bottom": 90, "end": 120},
    },
}


def _merge_profile(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    result = copy.deepcopy(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = _merge_profile(result[key], value)
        else:
            result[key] = value
    return result


def load_word_profile(path: Path | None = None) -> dict[str, Any]:
    if not path:
        return copy.deepcopy(DEFAULT_WORD_PROFILE)
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("word profile must be a JSON object")
    profile = _merge_profile(DEFAULT_WORD_PROFILE, data)
    if profile.get("profile_version") != "oks-word/v1":
        raise ValueError("word profile.profile_version must be oks-word/v1")
    table_width = profile["table"].get("width_twips")
    if not isinstance(table_width, int) or table_width <= 0:
        raise ValueError("word profile table.width_twips must be a positive integer")
    column_widths = profile["table"].get("column_widths")
    if not isinstance(column_widths, list) or not column_widths or any(float(value) <= 0 for value in column_widths):
        raise ValueError("word profile table.column_widths must be a non-empty list of positive numbers")
    return profile


def load_outline(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("outline must be a JSON object")
    if not str(data.get("title", "")).strip():
        raise ValueError("outline.title is required")
    if not isinstance(data.get("sections"), list):
        raise ValueError("outline.sections must be a list")
    if not isinstance(data.get("sources"), list) or not data["sources"]:
        raise ValueError("outline.sources must be a non-empty list")
    source_ids: set[str] = set()
    for index, source in enumerate(data["sources"]):
        if not isinstance(source, dict):
            raise ValueError(f"sources[{index}] must be an object")
        source_id = text(source.get("id"))
        label = text(source.get("label"))
        path_value = text(source.get("path"))
        status = text(source.get("status"))
        if not source_id or not label or not path_value:
            raise ValueError(f"sources[{index}] requires id, label, and path")
        if source_id in source_ids:
            raise ValueError(f"duplicate source id: {source_id}")
        if status not in VALID_SOURCE_STATUSES:
            raise ValueError(f"sources[{index}].status must be one of {sorted(VALID_SOURCE_STATUSES)}")
        source_ids.add(source_id)

    def check_refs(refs: Any, location: str) -> None:
        if not isinstance(refs, list) or not refs or any(not text(ref) for ref in refs):
            raise ValueError(f"{location}.source_refs must be a non-empty list")
        unknown = [text(ref) for ref in refs if text(ref) not in source_ids]
        if unknown:
            raise ValueError(f"{location}.source_refs contains unknown ids: {unknown}")

    if text(data.get("summary")):
        check_refs(data.get("summary_source_refs"), "summary")
    for section_index, section in enumerate(data["sections"]):
        location = f"sections[{section_index}]"
        if not isinstance(section, dict):
            raise ValueError(f"{location} must be an object")
        check_refs(section.get("source_refs"), location)
        for field in ("paragraphs", "bullets"):
            values = section.get(field, [])
            if not isinstance(values, list):
                raise ValueError(f"{location}.{field} must be a list")
            for item_index, item in enumerate(values):
                if isinstance(item, dict):
                    if not text(item.get("text")):
                        raise ValueError(f"{location}.{field}[{item_index}].text is required")
                    check_refs(item.get("source_refs"), f"{location}.{field}[{item_index}]")
                elif not text(item):
                    raise ValueError(f"{location}.{field}[{item_index}] cannot be empty")
        rows = section.get("table", [])
        if not isinstance(rows, list):
            raise ValueError(f"{location}.table must be a list")
        if rows:
            if any(not isinstance(row, list) or not row for row in rows):
                raise ValueError(f"{location}.table must contain non-empty row lists")
            width = len(rows[0])
            if any(len(row) != width for row in rows):
                raise ValueError(f"{location}.table must be rectangular")
    return data


@contextmanager
def atomic_target(output: Path):
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=f".{output.stem}-", suffix=output.suffix, dir=output.parent)
    os.close(fd)
    temp = Path(name)
    try:
        yield temp
        os.replace(temp, output)
    finally:
        temp.unlink(missing_ok=True)


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def item_text(item: Any) -> str:
    return text(item.get("text")) if isinstance(item, dict) else text(item)


def item_refs(item: Any, section: dict[str, Any]) -> list[str]:
    refs = item.get("source_refs") if isinstance(item, dict) else section.get("source_refs")
    return [text(ref) for ref in refs]


def source_suffix(item: Any, section: dict[str, Any]) -> str:
    refs = item_refs(item, section)
    return f" [sources: {', '.join(refs)}]" if refs else ""


def refs_suffix(refs: Any) -> str:
    values = [text(ref) for ref in refs] if isinstance(refs, list) else []
    return f" [sources: {', '.join(values)}]" if values else ""


def evidence_suffix(refs: Any, statuses: Any = None) -> str:
    source_part = refs_suffix(refs)
    values = [text(status) for status in statuses] if isinstance(statuses, list) else []
    status_part = f" [claim status: {', '.join(values)}]" if values else ""
    return source_part + status_part


def item_evidence_suffix(item: Any, section: dict[str, Any]) -> str:
    refs = item.get("source_refs") if isinstance(item, dict) else section.get("source_refs")
    statuses = item.get("claim_statuses") if isinstance(item, dict) else None
    return evidence_suffix(refs, statuses)


def table_blocks(section: dict[str, Any]) -> list[dict[str, Any]]:
    tables = section.get("tables")
    if isinstance(tables, list) and tables:
        return [table for table in tables if isinstance(table, dict)]
    legacy_rows = section.get("table", [])
    return [{"rows": legacy_rows, "source_refs": section.get("source_refs", [])}] if legacy_rows else []


def source_line(source: dict[str, Any]) -> str:
    details = [text(source.get("kind")), text(source.get("status"))]
    retrieved_at = text(source.get("retrieved_at"))
    if retrieved_at:
        details.append(f"retrieved {retrieved_at}")
    return f"{text(source.get('id'))} — {text(source.get('label'))}: {text(source.get('path'))} [{'; '.join(item for item in details if item)}]"


def _column_widths(rows: list[list[Any]], profile: dict[str, Any]) -> list[int]:
    count = len(rows[0])
    configured = [float(value) for value in profile["table"].get("column_widths", [])]
    if len(configured) == count and abs(sum(configured) - 1.0) < 0.02:
        weights = configured
    else:
        lengths = [max(1, max(len(text(row[index])) for row in rows)) for index in range(count)]
        weights = [max(0.12, min(0.58, length / sum(lengths))) for length in lengths]
        total = sum(weights)
        weights = [weight / total for weight in weights]
    total_twips = int(profile["table"]["width_twips"])
    widths = [int(total_twips * weight) for weight in weights[:-1]]
    widths.append(total_twips - sum(widths))
    return widths


def _set_docx_cell(cell: Any, width_twips: int, fill: str, profile: dict[str, Any], *, header: bool = False, align: Any = None) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_width = tc_pr.find(qn("w:tcW"))
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        tc_pr.append(tc_width)
    tc_width.set(qn("w:w"), str(width_twips))
    tc_width.set(qn("w:type"), "dxa")
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    margins = profile["table"]["cell_margin_twips"]
    cell_mar = tc_pr.find(qn("w:tcMar"))
    if cell_mar is None:
        cell_mar = OxmlElement("w:tcMar")
        tc_pr.append(cell_mar)
    for side in ("top", "start", "bottom", "end"):
        node = cell_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            cell_mar.append(node)
        node.set(qn("w:w"), str(int(margins[side])))
        node.set(qn("w:type"), "dxa")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if align is None:
        align = WD_ALIGN_PARAGRAPH.LEFT
    color = profile["colors"]["header_text" if header else "ink"]
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = profile["fonts"]["latin"]
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), profile["fonts"]["east_asia"])
            run.font.size = Pt(float(profile["fonts"]["table_pt"]))
            run.font.bold = header
            run.font.color.rgb = RGBColor.from_string(color)


def _style_docx_table(table: Any, rows: list[list[Any]], profile: dict[str, Any]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_width = tbl_pr.find(qn("w:tblW"))
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.append(tbl_width)
    tbl_width.set(qn("w:w"), str(int(profile["table"]["width_twips"])))
    tbl_width.set(qn("w:type"), "dxa")
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), profile["colors"]["border"])
    header_row = table.rows[0]._tr
    header_pr = header_row.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    header_pr.append(header_marker)
    widths = _column_widths(rows, profile)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            is_header = row_index == 0
            fill = profile["colors"]["accent"] if is_header else (profile["colors"]["row_alt"] if row_index % 2 == 0 else "FFFFFF")
            value = rows[row_index][col_index] if row_index < len(rows) and col_index < len(rows[row_index]) else ""
            numeric = not is_header and col_index > 0 and isinstance(value, (int, float))
            _set_docx_cell(cell, widths[col_index], fill, profile, header=is_header, align=WD_ALIGN_PARAGRAPH.RIGHT if numeric else WD_ALIGN_PARAGRAPH.LEFT)


def split_for_slide(value: str, limit: int = 360) -> list[str]:
    value = text(value)
    if len(value) <= limit:
        return [value]
    return [value[start : start + limit] for start in range(0, len(value), limit)]


def render_docx(outline: dict[str, Any], output: Path, profile: dict[str, Any] | None = None) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    profile = profile or load_word_profile()
    with atomic_target(output) as temp:
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(profile["page"]["top_twips"] / 1440)
        section.bottom_margin = Inches(profile["page"]["bottom_twips"] / 1440)
        section.left_margin = Inches(profile["page"]["left_twips"] / 1440)
        section.right_margin = Inches(profile["page"]["right_twips"] / 1440)
        normal = document.styles["Normal"]
        normal.font.name = profile["fonts"]["latin"]
        normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), profile["fonts"]["east_asia"])
        normal.font.size = Pt(float(profile["fonts"]["body_pt"]))
        for style_name, size, color in (("Title", 24, profile["colors"]["ink"]), ("Heading 1", 16, profile["colors"]["ink"]), ("Heading 2", 12, profile["colors"]["accent"])):
            style = document.styles[style_name]
            style.font.name = profile["fonts"]["latin"]
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), profile["fonts"]["east_asia"])
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string(color)

        title = document.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title.add_run(text(outline["title"]))
        if text(outline.get("subtitle")):
            document.add_paragraph(text(outline["subtitle"]))
        if text(outline.get("summary")):
            document.add_paragraph("Executive summary", style="Heading 1")
            document.add_paragraph(text(outline["summary"]) + evidence_suffix(outline.get("summary_source_refs"), outline.get("summary_claim_statuses")))
        for section_data in outline["sections"]:
            document.add_paragraph(text(section_data.get("heading", "Section")), style="Heading 1")
            for paragraph in section_data.get("paragraphs", []):
                document.add_paragraph(item_text(paragraph) + item_evidence_suffix(paragraph, section_data))
            for bullet in section_data.get("bullets", []):
                document.add_paragraph(item_text(bullet) + item_evidence_suffix(bullet, section_data), style="List Bullet")
            for table_data in table_blocks(section_data):
                rows = [row for row in table_data.get("rows", []) if isinstance(row, list) and row]
                if not rows:
                    continue
                note = document.add_paragraph(style="Caption")
                note.paragraph_format.space_before = Pt(5)
                note.paragraph_format.space_after = Pt(3)
                note_run = note.add_run("Evidence note · " + evidence_suffix(table_data.get("source_refs"), table_data.get("claim_statuses")).strip())
                note_run.font.name = profile["fonts"]["latin"]
                note_run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), profile["fonts"]["east_asia"])
                note_run.font.size = Pt(8)
                note_run.font.color.rgb = RGBColor.from_string(profile["colors"]["muted"])
                table = document.add_table(rows=1, cols=len(rows[0]))
                for cell, value in zip(table.rows[0].cells, rows[0]):
                    cell.text = text(value)
                for row in rows[1:]:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = text(value)
                _style_docx_table(table, rows, profile)
        document.add_paragraph("Sources", style="Heading 1")
        for source in outline["sources"]:
            document.add_paragraph(source_line(source), style="List Bullet")
        document.save(temp)


def _pdf_font(font_path: str | None):
    if not font_path:
        return "Helvetica", False
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    path = Path(font_path)
    if not path.is_file():
        raise ValueError(f"font path does not exist: {font_path}")
    pdfmetrics.registerFont(TTFont("OKSOfficeFont", str(path)))
    return "OKSOfficeFont", True


def contains_cjk(value: str) -> bool:
    return any(
        "\u2e80" <= char <= "\u9fff"
        or "\uf900" <= char <= "\ufaff"
        or "\uff00" <= char <= "\uffef"
        for char in value
    )


def render_pdf(outline: dict[str, Any], output: Path, font_path: str | None) -> None:
    if contains_cjk(json.dumps(outline, ensure_ascii=False)) and not font_path:
        raise ValueError("CJK content requires --font-path with a readable TTF font for PDF output")
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_name, custom_font = _pdf_font(font_path)
    if contains_cjk(json.dumps(outline, ensure_ascii=False)) and not custom_font:
        raise ValueError("CJK content requires --font-path with a readable TTF font for PDF output")
    with atomic_target(output) as temp:
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name="OKSTitle", parent=styles["Title"], fontName=font_name, fontSize=22, leading=27, textColor=colors.HexColor("#0f172a")))
        styles.add(ParagraphStyle(name="OKSHeading", parent=styles["Heading1"], fontName=font_name, fontSize=15, leading=19, textColor=colors.HexColor("#0f172a"), spaceBefore=12))
        styles.add(ParagraphStyle(name="OKSBody", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=15))
        doc = SimpleDocTemplate(str(temp), pagesize=A4, rightMargin=.7*inch, leftMargin=.7*inch, topMargin=.65*inch, bottomMargin=.65*inch)
        story = [Paragraph(html.escape(text(outline["title"])), styles["OKSTitle"]), Spacer(1, 12)]
        if text(outline.get("subtitle")):
            story += [Paragraph(html.escape(text(outline["subtitle"])), styles["OKSBody"]), Spacer(1, 8)]
        if text(outline.get("summary")):
            summary = html.escape(text(outline["summary"]) + evidence_suffix(outline.get("summary_source_refs"), outline.get("summary_claim_statuses")))
            story += [Paragraph("Executive summary", styles["OKSHeading"]), Paragraph(summary, styles["OKSBody"])]
        for section_data in outline["sections"]:
            story.append(Paragraph(html.escape(text(section_data.get("heading", "Section"))), styles["OKSHeading"]))
            for paragraph in section_data.get("paragraphs", []):
                story.append(Paragraph(html.escape(item_text(paragraph) + item_evidence_suffix(paragraph, section_data)), styles["OKSBody"]))
            for bullet in section_data.get("bullets", []):
                story.append(Paragraph("• " + html.escape(item_text(bullet) + item_evidence_suffix(bullet, section_data)), styles["OKSBody"]))
            for table_data in table_blocks(section_data):
                rows = [row for row in table_data.get("rows", []) if isinstance(row, list) and row]
                if not rows:
                    continue
                story.append(Paragraph(html.escape("Table" + evidence_suffix(table_data.get("source_refs"), table_data.get("claim_statuses"))), styles["OKSBody"]))
                table = Table([[Paragraph(html.escape(text(cell)), styles["OKSBody"]) for cell in row] for row in rows], repeatRows=1)
                table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbeafe")), ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#94a3b8")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("FONTNAME", (0, 0), (-1, -1), font_name), ("PADDING", (0, 0), (-1, -1), 6)]))
                story += [Spacer(1, 5), table]
        story.append(Paragraph("Sources", styles["OKSHeading"]))
        for source in outline["sources"]:
            story.append(Paragraph("• " + html.escape(source_line(source)), styles["OKSBody"]))
        doc.build(story)


def render_pptx(outline: dict[str, Any], output: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    with atomic_target(output) as temp:
        prs = Presentation()
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = text(outline["title"])
        if text(outline.get("subtitle")) and len(title_slide.placeholders) > 1:
            title_slide.placeholders[1].text = text(outline["subtitle"])
        if text(outline.get("summary")):
            summary_chunks = split_for_slide(text(outline["summary"]) + evidence_suffix(outline.get("summary_source_refs"), outline.get("summary_claim_statuses")))
            for slide_index in range(0, len(summary_chunks), 6):
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                suffix = "" if slide_index == 0 else " (continued)"
                slide.shapes.title.text = "Executive summary" + suffix
                body = slide.placeholders[1].text_frame
                body.clear()
                for item_index, item in enumerate(summary_chunks[slide_index : slide_index + 6]):
                    paragraph = body.paragraphs[0] if item_index == 0 else body.add_paragraph()
                    paragraph.text = item
                    paragraph.font.size = Pt(22)
        for section_data in outline["sections"]:
            items: list[str] = []
            for value in section_data.get("paragraphs", []) + section_data.get("bullets", []):
                items.extend(split_for_slide(item_text(value) + item_evidence_suffix(value, section_data)))
            for slide_index in range(0, max(len(items), 1), 6):
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                suffix = "" if slide_index == 0 else " (continued)"
                slide.shapes.title.text = text(section_data.get("heading", "Section")) + suffix
                body = slide.placeholders[1].text_frame
                body.clear()
                slide_items = items[slide_index : slide_index + 6]
                if not slide_items:
                    slide_items = ["No narrative content"]
                for index, item in enumerate(slide_items):
                    paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
                    paragraph.text = item
                    paragraph.level = 0
                    paragraph.font.size = Pt(20)
            for table_data in table_blocks(section_data):
                rows = [row for row in table_data.get("rows", []) if isinstance(row, list) and row]
                if not rows:
                    continue
                table_slide = prs.slides.add_slide(prs.slide_layouts[5])
                table_slide.shapes.title.text = text(section_data.get("heading", "Comparison")) + " — table" + evidence_suffix(table_data.get("source_refs"), table_data.get("claim_statuses"))
                shape = table_slide.shapes.add_table(len(rows), len(rows[0]), Inches(.5), Inches(1.4), Inches(12.3), Inches(5.3))
                table = shape.table
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        table.cell(r, c).text = text(value)
        source_slide = prs.slides.add_slide(prs.slide_layouts[1])
        source_slide.shapes.title.text = "Sources"
        body = source_slide.placeholders[1].text_frame
        body.clear()
        for index, source in enumerate(outline["sources"]):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = source_line(source)
            paragraph.font.size = Pt(16)
        prs.save(temp)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    input_group = parser.add_mutually_exclusive_group(required=True)
    input_group.add_argument("--outline", type=Path)
    input_group.add_argument("--package", type=Path, help="OKS Office evidence package; normalized to the shared outline contract")
    parser.add_argument("--format", required=True, choices=("docx", "pdf", "pptx", "ppt"))
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--font-path", help="TTF font for PDF output, required for reliable CJK glyphs")
    parser.add_argument("--word-profile", type=Path, help="OKS Word profile JSON for DOCX styling; defaults to the bundled OKS profile")
    args = parser.parse_args()
    outline = package_to_outline(load_package(args.package)) if args.package else load_outline(args.outline)
    fmt = "pptx" if args.format == "ppt" else args.format
    if fmt == "docx":
        render_docx(outline, args.output, load_word_profile(args.word_profile))
    elif fmt == "pdf":
        render_pdf(outline, args.output, args.font_path)
    else:
        render_pptx(outline, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
